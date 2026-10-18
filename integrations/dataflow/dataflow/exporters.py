"""Escritura de salidas: sin sobrescribir, neutralizando inyección CSV/fórmulas, con hash."""
from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path

from .errors import DataError
from .transforms import REASON

_DANGER = ("=", "+", "-", "@", "\t", "\r")
_NUMERIC = re.compile(r"^[+-]?\d+([.,]\d+)?$")


def neutralize(v):
    """Evita que una hoja de cálculo interprete texto como fórmula. Solo afecta a str."""
    if isinstance(v, str) and v and v[0] in _DANGER and not _NUMERIC.match(v):
        return "'" + v
    return v


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _out_rows(columns, rows):
    cols = list(columns) + ["_origen", "_fila", "_motivo"]
    for r in rows:
        yield cols, [neutralize(r.get(c, "")) for c in columns] + [
            r.get("__src", ""), r.get("__row", ""), neutralize(r.get(REASON, ""))]


def write_csv(path: Path, header: list[str], matrix: list[list]):
    with open(path, "x", encoding="utf-8", newline="") as f:  # 'x': falla si existe
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(matrix)
    return {"name": path.name, "sha256": _sha(path), "bytes": path.stat().st_size, "rows": len(matrix)}


def write_dataset_csv(path, columns, rows):
    items = list(_out_rows(columns, rows))
    header = items[0][0] if items else list(columns) + ["_origen", "_fila", "_motivo"]
    return write_csv(path, header, [m for _, m in items])


def write_xlsx(path: Path, sheets: dict):
    """sheets: nombre -> (header, matrix). Cadenas forzadas a tipo texto (sin fórmulas)."""
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        raise DataError("XLSX_UNAVAILABLE", "openpyxl no disponible para escribir xlsx") from None
    if path.exists():
        raise FileExistsError(path.name)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    total = 0
    for name, (header, matrix) in sheets.items():
        ws = wb.create_sheet(name[:31])
        ws.append(header)
        for row in matrix:
            ws.append(row)
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        total += len(matrix)
    wb.save(path)
    return {"name": path.name, "sha256": _sha(path), "bytes": path.stat().st_size, "rows": total}
